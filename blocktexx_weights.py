"""Private historical pickup imports; source records live only in the saved model."""
import hashlib
import io
import re
from datetime import datetime
from zipfile import ZipFile, BadZipFile


def normalized(value):
    return re.sub(r'[^a-z0-9]', '', value.lower().replace('blocktexx', '').replace('pty ltd', ''))


def validate_history(raw, sites, number, text):
    if not isinstance(raw, dict):
        raise ValueError('Invalid pickup history.')
    rows = raw.get('pickups', [])
    if not isinstance(rows, list) or len(rows) > 2000:
        raise ValueError('Use at most 2,000 historical pickups.')
    ids = {s['id'] for s in sites}
    pickups = []
    seen = set()
    for row in rows:
        if not isinstance(row, dict):
            raise ValueError('Invalid historical pickup.')
        item = {k: text(row.get(k, ''), k, 250) for k in ('date','docket','company','site','facility','profile','site_id')}
        try:
            datetime.strptime(item['date'], '%Y-%m-%d')
        except ValueError:
            raise ValueError('Historical pickups need an ISO date.') from None
        key = tuple(item[k] for k in ('date','docket','company','site','facility'))
        if key in seen:
            raise ValueError('Duplicate pickup docket: combine its material lines first.')
        seen.add(key)
        if not item['profile'] or not item['docket'] or (item['site_id'] and item['site_id'] not in ids):
            raise ValueError('Invalid historical pickup mapping.')
        item['kg'] = number(row.get('kg'), 'Historical kg', 1000000)
        pickups.append(item)
    return {'source': text(raw.get('source',''), 'History source', 250), 'pickups': pickups}


def import_workbook(content, filename, sites):
    from openpyxl import load_workbook
    if len(content) > 5_000_000:
        raise ValueError('Use a workbook smaller than 5 MB.')
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 40_000_000:
                raise ValueError('Workbook expands beyond 40 MB.')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError):
        raise ValueError('Upload a valid Excel .xlsx workbook.') from None
    if 'KG Collected' not in workbook.sheetnames:
        raise ValueError('Workbook needs a KG Collected sheet.')
    rows = workbook['KG Collected'].iter_rows(values_only=True)
    headers = None
    grouped = {}
    for index, row in enumerate(rows):
        if index > 10000:
            raise ValueError('History sheet exceeds 10,000 rows.')
        if headers is None:
            if 'Docket' in row and 'Qty' in row:
                headers = {str(v).strip(): i for i,v in enumerate(row) if v is not None}
                if not {'Date','Docket','Company','Site','Facility','Qty'} <= headers.keys():
                    raise ValueError('Missing history columns.')
            continue
        get = lambda name: row[headers[name]] if len(row)>headers[name] else None
        if not get('Docket'):
            continue
        date = get('Date')
        if isinstance(date, datetime):
            date = date.strftime('%Y-%m-%d')
        else:
            try:
                date = datetime.strptime(str(date), '%m/%d/%Y').strftime('%Y-%m-%d')
            except ValueError:
                raise ValueError('Unrecognized pickup date on row '+str(index+1)) from None
        key = (date, str(get('Docket')), str(get('Company') or ''), str(get('Site') or ''), str(get('Facility') or ''))
        qty = get('Qty')
        if isinstance(qty, bool) or not isinstance(qty, (int,float)) or qty < 0:
            raise ValueError('Invalid kg on history row '+str(index+1))
        grouped[key] = grouped.get(key, 0) + qty
    workbook.close()
    if not grouped:
        raise ValueError('No pickup weights found.')
    pickups = []
    for key, kg in grouped.items():
        date,docket,company,site,facility = key
        profile = hashlib.sha256('|'.join((company,site,facility)).encode()).hexdigest()[:16]
        state_match = re.match(r'SX(NSW|QLD|VIC|SA|ACT|WA|TAS)',facility)
        state = state_match.group(1) if state_match else ''
        eligible = [s for s in sites if s['state']==state]
        full = normalized(company+site)
        exact = [s for s in eligible if normalized(s['name'])==full]
        # A unique company is sufficient only if no competing site exists in that state.
        company_matches = [s for s in eligible if normalized(company) in normalized(s['name'])]
        candidates = exact or (company_matches if len(company_matches)==1 else [])
        site_id = candidates[0]['id'] if len(candidates)==1 else ''
        pickups.append(dict(date=date,docket=docket,company=company,site=site,facility=facility,profile=profile,kg=kg,site_id=site_id))
    return {'source': filename, 'pickups': pickups}
